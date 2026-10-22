"""Online Supplement Section O1: summarize partials/ (see prominence_sensitivity.py).

Writes
  TableO1_prominence_JohnsonSU.csv     Table O1 -- false modality (%) for Johnson SU,
                                       N = 50 and 200, K = 7, 10, 13, at 0.5%, 1%, 2%
  excel/Prominence_Sensitivity.xlsx   workbook covering every section, all thresholds
  prominence_sensitivity_summary.csv   the workbook's per-section tables in one file
"""
import os, glob, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
TH = [0.005, 0.01, 0.02, 0.05]; TL = {0.005: '0.5%', 0.01: '1% (paper)', 0.02: '2%', 0.05: '5%'}
ORDER = ['GEV', 'Metalog', 'Log Metalog', 'QFlex-U', 'Log QFlex-U', 'QFlex-T+', 'Log QFlex-T+',
         'QFlex-TA+', 'Log QFlex-TA+', 'QFlex-A+', 'Log QFlex-A+']
frames = []
for f in sorted(glob.glob(os.path.join(HERE, 'partials', '*.csv'))):
    d = pd.read_csv(f, keep_default_na=False); d['valid'] = d['valid'].astype(str) == 'True'; frames.append(d)
d = pd.concat(frames, ignore_index=True)
for t in TH:
    d[f'm{t}'] = [np.nan if not v or r == 'NA' else max(1, sum(float(x) >= t for x in r.split(';') if x))
                  for v, r in zip(d.valid, d.ratios)]
SHEETS = [('MC Johnson', lambda s: s.startswith('mc_'), 'False multimodality'),
          ('Inverse Gaussian', lambda s: s.startswith('ig_'), 'False multimodality'),
          ('Bimodal mixture', lambda s: s == 'bimodal', 'Multimodal'),
          ('SU bootstrap', lambda s: s == 'bootstrap', 'False multimodality'),
          ('River gauge', lambda s: s == 'hydrology', 'Multimodal'),
          ('Fish weights', lambda s: s.startswith('fish_'), 'Multimodal'),
          ('Geyser', lambda s: s == 'geyser', 'Multimodal')]
def table(sub):
    rows = []
    for (st, mo, K), g in sub.groupby(['setting', 'model', 'K']):
        v = g[g.valid]; r = dict(setting=st, model=mo, K=K, total=len(g), valid=len(v))
        for t in TH:
            mm = v[f'm{t}'].dropna(); n = len(mm)
            for k, lab in ((1, '1'), (2, '2'), (3, '3')):
                r[f'{t}_{lab}'] = 100 * (mm == k).mean() if n else np.nan
            r[f'{t}_4+'] = 100 * (mm >= 4).mean() if n else np.nan
            r[f'{t}_multi'] = 100 * (mm > 1).mean() if n else np.nan
        rows.append(r)
    t = pd.DataFrame(rows)
    t['mo'] = t.model.map({m: i for i, m in enumerate(ORDER)})
    # settings: Johnson SU, SL, SB; then N = 50 before N = 200 (fish: jitter 0.5 before 1.0)
    dist_rank = lambda st: next((i for i, k in enumerate(('Johnson SU', 'Johnson SL', 'Johnson SB')) if st.startswith(k)), 9)
    n_of = lambda st: int(st.split('N=')[1]) if 'N=' in st else 0
    t['d_'] = t.setting.map(dist_rank); t['n_'] = t.setting.map(n_of)
    return t.sort_values(['d_', 'n_', 'setting', 'mo', 'K']).drop(columns=['mo', 'd_', 'n_'])
F = lambda **k: Font(name='Arial', **k)
thin = Side(style='thin', color='BFBFBF'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
HDR = PatternFill('solid', fgColor='D9E1F2'); SUB = PatternFill('solid', fgColor='F2F2F2'); PAPER = PatternFill('solid', fgColor='FFF2CC')
wb = Workbook(); ws0 = wb.active; ws0.title = 'Read me'
notes = ['Sensitivity of mode identification to the relative-prominence threshold',
 '',
 'Rule: an interior local maximum of the grid-difference density counts as a mode when its prominence >= t x its own height.',
 'The paper uses t = 1%. Here t = 0.5%, 1%, 2% and 5%. A fit with no qualifying peak is unimodal (1 mode).',
 'Every fit was redone with the section\'s own sampler and fitting code; for each valid fit the relative prominence of every',
 'local maximum was stored, so all thresholds are evaluated on the SAME fits (only the counting rule changes).',
 'Check: at t = 1% the per-fit validity and mode counts reproduce the paper\'s stored results exactly (MC Johnson, inverse Gaussian,',
 'bimodal, SU bootstrap, river gauge, geyser: fit-by-fit; fish: identical distributions of mode counts per model and K).',
 '',
 'Columns: "valid" = valid fits out of "total". Mode-count shares (1, 2, 3, 4+) and Multimodal % (= more than one mode) are % of VALID fits.',
 'For unimodal truths (Johnson SU/SL/SB, inverse Gaussian, the single SU realization) multimodal = false multimodality.',
 'For the bimodal mixture the true count is 2. River gauge, fish and geyser are real data (true count unknown).',
 'The 1% (paper) columns are shaded.',
 '',
 'Settings (as in the paper): MC N = 50 and 200, 1000 replicates, K = 4/7/10/13; inverse Gaussian N = 200, same design; bimodal',
 'mixture N = 200, 1000 replicates, K = 4..14; SU bootstrap 1000 resamples of the N = 200 realization; river gauge, fish (jitter',
 '0.5 and 1.0 lb) and geyser (jitter 0.5 min): 1000 bootstrap replicates each. Each section uses the QPDs the paper reports for it:',
 'all five (including QFlex-T+) for the Monte Carlo, inverse Gaussian and SU bootstrap; the section\'s own models elsewhere.',
 'Source: reproduction/prominence_sensitivity/ (prominence_sensitivity.py, summarize.py, partials/).']
for i, t in enumerate(notes, 1):
    ws0.cell(i, 1, t).font = F(bold=(i == 1), size=13 if i == 1 else 10)
ws0.column_dimensions['A'].width = 140
summ = []
for name, sel, mlab in SHEETS:
    t = table(d[d.section.map(sel)]); summ.append((name, t))
    ws = wb.create_sheet(name)
    ws.cell(1, 1, f'{name}: mode counts by prominence threshold (% of valid fits)').font = F(bold=True, size=12)
    base = ['Setting', 'Model', 'K', 'Total fits', 'Valid fits']
    sub = ['1 mode', '2 modes', '3 modes', '4+ modes', f'{mlab} %']
    r1, r2 = 3, 4
    for j, h in enumerate(base, 1):
        c = ws.cell(r1, j, h); ws.merge_cells(start_row=r1, start_column=j, end_row=r2, end_column=j)
        for rr in (r1, r2): ws.cell(rr, j).fill = HDR; ws.cell(rr, j).border = B
        c.font = F(bold=True); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    col = len(base) + 1
    for t_ in TH:
        c = ws.cell(r1, col, f'Prominence threshold {TL[t_]}'); ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + 4)
        for k in range(5):
            x = ws.cell(r1, col + k); x.fill = PAPER if t_ == 0.01 else HDR; x.border = B
            y = ws.cell(r2, col + k, sub[k]); y.font = F(bold=True); y.fill = PAPER if t_ == 0.01 else SUB; y.border = B
            y.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.font = F(bold=True); c.alignment = Alignment(horizontal='center')
        col += 5
    ws.row_dimensions[r2].height = 30
    r = r2 + 1; prev = None
    for _, row in t.iterrows():
        vals = [row.setting if row.setting != prev else None, row.model, ('-' if row.K == 0 else int(row.K)), int(row.total), int(row.valid)]
        prev = row.setting
        for t_ in TH:
            vals += [row[f'{t_}_1'], row[f'{t_}_2'], row[f'{t_}_3'], row[f'{t_}_4+'], row[f'{t_}_multi']]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, None if (isinstance(v, float) and np.isnan(v)) else v); c.font = F(bold=(j == 1)); c.border = B
            if j > 5:
                c.number_format = '0.0'; c.alignment = Alignment(horizontal='center')
                if (j - 6) // 5 == 1: c.fill = PAPER
                if (j - 6) % 5 == 4: c.font = F(bold=True)
        r += 1
    widths = [30, 15, 5, 8, 8] + [8, 8, 8, 8, 13] * 4
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r2 + 1, 6)
# compact summary: multimodal % only
ws = wb.create_sheet('Summary', 1)
ws.cell(1, 1, 'Multimodal % of valid fits (false multimodality where the truth is unimodal) at each prominence threshold').font = F(bold=True, size=12)
hdr = ['Section', 'Setting', 'Model', 'K', 'Valid fits'] + [TL[t_] for t_ in TH] + ['Max change vs 1% (pp)']
for j, h in enumerate(hdr, 1):
    c = ws.cell(3, j, h); c.font = F(bold=True); c.fill = PAPER if h == TL[0.01] else HDR; c.border = B
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
r = 4
for name, t in summ:
    for _, row in t.iterrows():
        mv = [row[f'{t_}_multi'] for t_ in TH]
        ch = np.nanmax(np.abs(np.array(mv) - row['0.01_multi'])) if row.valid else np.nan
        vals = [name, row.setting, row.model, ('-' if row.K == 0 else int(row.K)), int(row.valid)] + mv + [ch]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, None if (isinstance(v, float) and np.isnan(v)) else v); c.font = F(); c.border = B
            if j > 5: c.number_format = '0.0'; c.alignment = Alignment(horizontal='center')
            if j == 7: c.fill = PAPER
        r += 1
for j, w in enumerate([16, 34, 15, 5, 8, 9, 11, 9, 9, 12], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'F4'
out = os.path.abspath(os.path.join(HERE, '..', '..', 'excel', 'Prominence_Sensitivity.xlsx')); wb.save(out)
allt = pd.concat([t.assign(section=n) for n, t in summ]); allt.to_csv(os.path.join(HERE, 'prominence_sensitivity_summary.csv'), index=False)

# Online Supplement Table O1 (Section O1): false modality (%) for Johnson SU,
# N = 50 and 200, K = 7, 10, 13, at prominence thresholds 0.5%, 1% and 2%.
_su = allt[allt.setting.str.match(r'Johnson SU, N=\d+$') & allt.K.isin([7, 10, 13])].copy()
_su['N'] = _su.setting.str.extract(r'N=(\d+)')[0].astype(int)
_o1 = _su[['N', 'model', 'K'] + [f'{t}_multi' for t in (0.005, 0.01, 0.02)]].rename(
    columns={'0.005_multi': 'FalseModality_pct_tau0.5', '0.01_multi': 'FalseModality_pct_tau1.0',
             '0.02_multi': 'FalseModality_pct_tau2.0'})
_o1.round(1).to_csv(os.path.join(HERE, 'TableO1_prominence_JohnsonSU.csv'), index=False)
print('saved', out)


# ------------------------------------------------------------------ workbook companion sheet
# False multimodality (% of valid fits) at prominence 0.5%, 1% (paper) and 2%, for each
# Johnson distribution x N x QPD, at K = 7, 10, 13; plus the largest change from 1%.
# Its Johnson SU rows are Table O1; the SL and SB rows are not in the paper.
from openpyxl import load_workbook
mc = dict(summ)['MC Johnson']
wb = load_workbook(out)
ws = wb.create_sheet('Companion table', 1)
TC = [0.005, 0.01, 0.02]; KC = [7, 10, 13]
MODELS_C = ['Metalog', 'QFlex-U', 'QFlex-T+', 'QFlex-TA+', 'QFlex-A+']
ws.cell(1, 1, 'False multimodality rate (% of valid fits) by relative-prominence threshold, '
              'Monte Carlo realizations of the Johnson distributions (1000 replicates per N). '
              'The Johnson SU rows are Online Supplement Table O1.').font = F(bold=True, size=11)
hdr1 = ['Johnson dist.', 'N', 'QPD']
for j, h in enumerate(hdr1, 1):
    c = ws.cell(3, j, h); ws.merge_cells(start_row=3, start_column=j, end_row=4, end_column=j)
    for rr in (3, 4): ws.cell(rr, j).fill = HDR; ws.cell(rr, j).border = B
    c.font = F(bold=True); c.alignment = Alignment(horizontal='center', vertical='center')
col = 4
for K in KC:
    c = ws.cell(3, col, f'K = {K}'); ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 2)
    c.font = F(bold=True); c.alignment = Alignment(horizontal='center')
    for k, t_ in enumerate(TC):
        ws.cell(3, col + k).fill = HDR; ws.cell(3, col + k).border = B
        y = ws.cell(4, col + k, {0.005: '0.5%', 0.01: '1%', 0.02: '2%'}[t_]); y.font = F(bold=True)
        y.fill = PAPER if t_ == 0.01 else SUB; y.border = B; y.alignment = Alignment(horizontal='center')
    col += 3
c = ws.cell(3, col, 'Max change from 1% (pp)'); ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
for rr in (3, 4): ws.cell(rr, col).fill = HDR; ws.cell(rr, col).border = B
c.font = F(bold=True); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
r = 5
for dist in ('SU', 'SL', 'SB'):
    for n in (50, 200):
        st = f'Johnson {dist}, N={n}'
        for i, mo in enumerate(MODELS_C):
            vals = [dist if (n == 50 and i == 0) else None, n if i == 0 else None, mo]
            ch = 0.0
            for K in KC:
                row = mc[(mc.setting == st) & (mc.model == mo) & (mc.K == K)].iloc[0]
                v = [row[f'{t_}_multi'] for t_ in TC]
                vals += v
                if row.valid: ch = max(ch, abs(v[0] - v[1]), abs(v[2] - v[1]))
            vals.append(ch)
            for j, v in enumerate(vals, 1):
                cc = ws.cell(r, j, None if (isinstance(v, float) and np.isnan(v)) else v)
                cc.font = F(bold=(j == 1)); cc.border = B
                if j >= 4:
                    cc.number_format = '0.0'; cc.alignment = Alignment(horizontal='center')
                    if j < col and (j - 4) % 3 == 1: cc.fill = PAPER
                elif j == 2: cc.alignment = Alignment(horizontal='center')
            r += 1
r += 1
for t in ['Rule: a local maximum of the fitted density is a mode when its prominence is at least the threshold times its own height; '
          'the paper uses 1% (shaded).',
          'False multimodality = share of valid fits with more than one mode. All thresholds are applied to the same fits; '
          'valid-fit counts are in the MC Johnson sheet.',
          'Fits: Metalog and QFlex-U (unconstrained), QFlex-T+, QFlex-TA+ and QFlex-A+; Log/Logit transforms with the true bounds '
          'for Johnson SL/SB; evaluation framework of Section 3.']:
    ws.cell(r, 1, t).font = F(italic=True, size=9, color='595959'); r += 1
for j, w in enumerate([13, 6, 12] + [8] * 9 + [12], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[3].height = 30
ws.freeze_panes = 'D5'
wb.save(out); print('companion table added')
